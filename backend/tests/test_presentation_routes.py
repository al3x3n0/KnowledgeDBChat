"""The presentation routes, run against a database with storage faked.

There were no tests of this endpoint. These pin what a user sees, and three
things that were wrong: a PPTX template could not be uploaded at all, a deck
whose title is not in Latin script could not be downloaded, and a negative
page size returned every job.
"""

from __future__ import annotations

import io
from uuid import uuid4

import pytest
from fastapi import HTTPException

from app.api.endpoints import presentations
from app.models.presentation import PresentationJob, PresentationTemplate
from app.schemas.presentation import (
    PresentationJobCreate,
    PresentationTemplateCreate,
    PresentationTemplateUpdate,
)
from app.services.storage_service import StorageService

PPTX = "application/vnd.openxmlformats-officedocument.presentationml.presentation"


@pytest.fixture
def storage(monkeypatch):
    """Object storage as a dict, and what was asked of it."""
    files, removed = {}, []

    async def upload_to_path(self, object_path, content, content_type=None):
        files[object_path] = content
        return object_path

    async def get_file_content(self, object_path):
        if object_path not in files:
            raise FileNotFoundError(object_path)
        return files[object_path]

    async def delete_file(self, object_path):
        removed.append(object_path)
        return files.pop(object_path, None) is not None

    monkeypatch.setattr(StorageService, "upload_to_path", upload_to_path)
    monkeypatch.setattr(StorageService, "get_file_content", get_file_content)
    monkeypatch.setattr(StorageService, "delete_file", delete_file)

    class Storage:
        pass

    fake = Storage()
    fake.files, fake.removed = files, removed
    return fake


@pytest.fixture
def queued(monkeypatch):
    from app.tasks.presentation_tasks import generate_presentation_task

    sent = []
    monkeypatch.setattr(
        generate_presentation_task, "delay", lambda job_id, user_id: sent.append(job_id)
    )
    return sent


def _deck() -> bytes:
    from pptx import Presentation

    deck, out = Presentation(), io.BytesIO()
    deck.slides.add_slide(deck.slide_layouts[0])
    deck.save(out)
    return out.getvalue()


class _Upload:
    def __init__(self, filename, content, content_type=PPTX):
        self.filename, self._content, self.content_type = (
            filename,
            content,
            content_type,
        )

    async def read(self):
        return self._content


async def _upload(db, user, filename="Corporate Theme.pptx", content=None):
    return await presentations.upload_pptx_template(
        file=_Upload(filename, _deck() if content is None else content),
        name="Corporate",
        description="house style",
        is_public=False,
        current_user=user,
        db=db,
    )


async def _job(db, user, **fields):
    job = PresentationJob(
        id=uuid4(),
        user_id=user.id if hasattr(user, "id") else user,
        title=fields.pop("title", "Quarterly"),
        topic="results",
        source_document_ids=[],
        slide_count=10,
        style="professional",
        include_diagrams=1,
        status=fields.pop("status", "completed"),
        progress=100,
        **fields,
    )
    db.add(job)
    await db.commit()
    return job


class TestUploadingATemplate:
    @pytest.mark.asyncio
    async def test_the_deck_is_stored_where_the_template_says_it_is(
        self, db_session, test_user, storage
    ):
        # It could not be uploaded at all: the storage call was given an
        # argument that function does not take.
        answer = await _upload(db_session, test_user)

        template = await db_session.get(PresentationTemplate, answer.id)
        assert template.template_type == "pptx"
        assert template.file_path == f"templates/{template.id}/Corporate_Theme.pptx"
        assert list(storage.files) == [template.file_path]
        assert storage.files[template.file_path][:2] == b"PK"

    @pytest.mark.asyncio
    async def test_a_file_that_is_not_a_deck(self, db_session, test_user, storage):
        with pytest.raises(HTTPException) as wrong_name:
            await _upload(db_session, test_user, filename="notes.txt")
        with pytest.raises(HTTPException) as not_a_deck:
            await _upload(db_session, test_user, content=b"not a zip at all")

        assert wrong_name.value.status_code == 400
        assert not_a_deck.value.status_code == 400
        assert storage.files == {}

    @pytest.mark.asyncio
    async def test_deleting_the_template_removes_its_file(
        self, db_session, test_user, storage
    ):
        answer = await _upload(db_session, test_user)
        path = next(iter(storage.files))

        await presentations.delete_template(
            template_id=answer.id, current_user=test_user, db=db_session
        )

        assert storage.removed == [path] and storage.files == {}
        assert await db_session.get(PresentationTemplate, answer.id) is None


class TestTemplates:
    async def _template(self, db, owner_id, **fields):
        template = PresentationTemplate(
            id=uuid4(),
            user_id=owner_id,
            name=fields.pop("name", "Theme"),
            template_type="theme",
            is_public=fields.pop("is_public", False),
            is_system=fields.pop("is_system", False),
            **fields,
        )
        db.add(template)
        await db.commit()
        return template

    @pytest.mark.asyncio
    async def test_you_see_your_own_and_the_shared_ones(self, db_session, test_user):
        other = uuid4()
        await self._template(db_session, test_user.id, name="mine")
        await self._template(db_session, other, name="public", is_public=True)
        await self._template(db_session, other, name="system", is_system=True)
        await self._template(db_session, other, name="private")
        await self._template(db_session, test_user.id, name="retired", is_active=False)

        everything = await presentations.list_templates(
            include_system=True,
            include_public=True,
            current_user=test_user,
            db=db_session,
        )
        own = await presentations.list_templates(
            include_system=False,
            include_public=False,
            current_user=test_user,
            db=db_session,
        )

        assert [t.name for t in everything] == ["system", "mine", "public"]
        assert [t.name for t in own] == ["mine"]

    @pytest.mark.asyncio
    async def test_somebody_elses_private_template_is_not_found(
        self, db_session, test_user
    ):
        private = await self._template(db_session, uuid4())

        with pytest.raises(HTTPException) as refused:
            await presentations.get_template(
                template_id=private.id, current_user=test_user, db=db_session
            )

        assert refused.value.status_code == 404

    @pytest.mark.asyncio
    async def test_only_the_owner_changes_a_template(self, db_session, test_user):
        public = await self._template(db_session, uuid4(), is_public=True)
        mine = await self._template(db_session, test_user.id)

        with pytest.raises(HTTPException) as refused:
            await presentations.update_template(
                template_id=public.id,
                request=PresentationTemplateUpdate(name="taken"),
                current_user=test_user,
                db=db_session,
            )
        changed = await presentations.update_template(
            template_id=mine.id,
            request=PresentationTemplateUpdate(name="renamed", is_public=True),
            current_user=test_user,
            db=db_session,
        )

        assert refused.value.status_code == 404
        assert (changed.name, changed.is_public) == ("renamed", True)

    @pytest.mark.asyncio
    async def test_a_system_template_is_nobodys_to_change(self, db_session, test_user):
        system = await self._template(db_session, test_user.id, is_system=True)

        with pytest.raises(HTTPException) as update:
            await presentations.update_template(
                template_id=system.id,
                request=PresentationTemplateUpdate(name="x"),
                current_user=test_user,
                db=db_session,
            )
        with pytest.raises(HTTPException) as delete:
            await presentations.delete_template(
                template_id=system.id, current_user=test_user, db=db_session
            )

        assert (update.value.status_code, delete.value.status_code) == (403, 403)

    @pytest.mark.asyncio
    async def test_creating_a_theme(self, db_session, test_user):
        made = await presentations.create_template(
            request=PresentationTemplateCreate(name="Dark", template_type="theme"),
            current_user=test_user,
            db=db_session,
        )

        assert (made.name, made.is_system, made.user_id) == (
            "Dark",
            False,
            test_user.id,
        )


class TestJobs:
    @pytest.mark.asyncio
    async def test_creating_one_queues_it(self, db_session, test_user, queued):
        made = await presentations.create_presentation(
            request=PresentationJobCreate(title="Q3", topic="results", slide_count=8),
            current_user=test_user,
            db=db_session,
        )

        assert (made.status, made.progress, made.slide_count) == ("pending", 0, 8)
        assert queued == [str(made.id)]

    @pytest.mark.asyncio
    async def test_a_template_you_may_not_use(self, db_session, test_user, queued):
        private = PresentationTemplate(
            id=uuid4(), user_id=uuid4(), name="theirs", template_type="theme"
        )
        db_session.add(private)
        await db_session.commit()

        with pytest.raises(HTTPException) as refused:
            await presentations.create_presentation(
                request=PresentationJobCreate(
                    title="Q3", topic="results", template_id=private.id
                ),
                current_user=test_user,
                db=db_session,
            )

        assert refused.value.status_code == 403
        assert queued == []

    @pytest.mark.asyncio
    async def test_listing_is_yours_newest_first_and_filtered(
        self, db_session, test_user
    ):
        from datetime import datetime, timedelta

        for n, status in enumerate(["completed", "failed", "completed"]):
            await _job(
                db_session,
                test_user,
                title=f"deck {n}",
                status=status,
                created_at=datetime(2026, 9, 1) + timedelta(minutes=n),
            )
        await _job(db_session, uuid4(), title="somebody else's")

        everything = await presentations.list_presentations(
            limit=20,
            offset=0,
            status_filter=None,
            current_user=test_user,
            db=db_session,
        )
        done = await presentations.list_presentations(
            limit=1,
            offset=1,
            status_filter="completed",
            current_user=test_user,
            db=db_session,
        )

        assert [job.title for job in everything] == ["deck 2", "deck 1", "deck 0"]
        assert [job.title for job in done] == ["deck 0"]

    def test_a_page_size_is_bounded(self):
        # A negative limit is "no limit" to the database.
        import inspect

        parameters = inspect.signature(presentations.list_presentations).parameters
        limit, offset = parameters["limit"].default, parameters["offset"].default
        bounds = {type(m).__name__: m for m in limit.metadata}
        offset_bounds = {type(m).__name__: m for m in offset.metadata}

        assert bounds["Ge"].ge == 1 and bounds["Le"].le == 100
        assert offset_bounds["Ge"].ge == 0

    @pytest.mark.asyncio
    async def test_somebody_elses_job_is_not_found(self, db_session, test_user):
        theirs = await _job(db_session, uuid4())

        for call in (
            presentations.get_presentation_job,
            presentations.download_presentation,
            presentations.delete_presentation,
            presentations.cancel_presentation,
        ):
            with pytest.raises(HTTPException) as refused:
                await call(job_id=theirs.id, current_user=test_user, db=db_session)
            assert refused.value.status_code == 404

    @pytest.mark.asyncio
    @pytest.mark.parametrize(
        "title", ["Quarterly", "Отчёт за квартал", 'So-called "plan"']
    )
    async def test_downloading_a_deck_whatever_its_title(
        self, db_session, test_user, storage, title
    ):
        # A title in another script made the response itself raise.
        job = await _job(
            db_session, test_user, title=title, file_path=f"presentations/x/{title}"
        )
        storage.files[job.file_path] = b"PK-deck"

        answer = await presentations.download_presentation(
            job_id=job.id, current_user=test_user, db=db_session
        )

        assert answer.status_code == 200 and answer.body == b"PK-deck"
        assert answer.headers["content-disposition"].startswith("attachment; filename=")

    @pytest.mark.asyncio
    async def test_a_deck_that_is_not_ready_or_not_there(
        self, db_session, test_user, storage
    ):
        running = await _job(db_session, test_user, status="generating")
        no_file = await _job(db_session, test_user)
        lost = await _job(db_session, test_user, file_path="presentations/gone.pptx")

        codes = []
        for job in (running, no_file, lost):
            with pytest.raises(HTTPException) as refused:
                await presentations.download_presentation(
                    job_id=job.id, current_user=test_user, db=db_session
                )
            codes.append(refused.value.status_code)

        assert codes == [400, 404, 404]

    @pytest.mark.asyncio
    async def test_deleting_removes_the_deck_too(self, db_session, test_user, storage):
        job = await _job(db_session, test_user, file_path="presentations/x/deck.pptx")
        storage.files[job.file_path] = b"PK"

        await presentations.delete_presentation(
            job_id=job.id, current_user=test_user, db=db_session
        )

        assert storage.removed == ["presentations/x/deck.pptx"]
        assert await db_session.get(PresentationJob, job.id) is None

    @pytest.mark.asyncio
    async def test_cancelling(self, db_session, test_user):
        running = await _job(db_session, test_user, status="generating")
        finished = await _job(db_session, test_user, status="completed")

        cancelled = await presentations.cancel_presentation(
            job_id=running.id, current_user=test_user, db=db_session
        )
        with pytest.raises(HTTPException) as refused:
            await presentations.cancel_presentation(
                job_id=finished.id, current_user=test_user, db=db_session
            )

        assert (cancelled.status, cancelled.error) == ("cancelled", "Cancelled by user")
        assert refused.value.status_code == 400
