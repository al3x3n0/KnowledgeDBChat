"""The training, template-fill and repository-report Celery tasks, run for real.

No test named any of these tasks, so nothing had executed their bodies. These
tests run them -- the async implementations, and the synchronous Celery
functions via ``.run`` -- against real rows, and judge on the rows they leave.

Only edges that leave the process are replaced, and every replacement binds
its arguments against the real callee's signature, so a call the real thing
would refuse fails here too:

* Redis -- ``job_support.publish_message`` / ``publish_progress`` /
  ``publish_sync`` / ``flag_is_set`` / ``delete_keys``, and the
  ``redis.asyncio`` client ``TrainingService.cancel_job`` opens;
* MinIO -- ``MinIOStorageService`` (the class of both the ``storage_service``
  singleton and every ``StorageService()``);
* the model -- ``LLMService.generate_response``;
* the vector store -- ``VectorStoreService.initialize`` / ``search``;
* the network -- every ``httpx.AsyncClient`` gets a ``MockTransport`` that
  answers as the GitHub API (and as a diagram renderer).

The trainer is the real ``simulated`` backend; the DOCX, PDF and PPTX builders
are real and their bytes are opened again with python-docx and python-pptx.

A test marked xfail(strict) is a defect in the application, not in the test.
"""

import asyncio
import base64
import inspect
import io
import json
import os
from uuid import UUID, uuid4

import httpx
import pytest
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession, async_sessionmaker, create_async_engine
from sqlalchemy.pool import NullPool

from app.core.config import settings
from app.core.database import Base
from app.models.document import Document, DocumentSource
from app.models.model_registry import AdapterStatus, ModelAdapter
from app.models.repo_report import RepoReportJob
from app.models.template import TemplateJob
from app.models.training_dataset import DatasetSample, TrainingDataset
from app.models.training_job import TrainingCheckpoint, TrainingJob
from app.services.auth_service import AuthService
from app.services.llm_service import LLMService
from app.services.storage_service import MinIOStorageService
from app.services.training_service import training_service
from app.services.vector_store import VectorStoreService
from app.tasks import job_support, repo_report_tasks, template_tasks, training_tasks

pytestmark = pytest.mark.unit

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
PPTX_MIME = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
TASK_MODULES = (training_tasks, template_tasks, repo_report_tasks, job_support)


# --------------------------------------------------------------------------
# Edges that leave the process
# --------------------------------------------------------------------------


def _binding(real):
    """A checker that raises TypeError exactly when ``real`` would."""
    signature = inspect.signature(real)

    def bind(*args, **kwargs):
        bound = signature.bind(*args, **kwargs)
        bound.apply_defaults()
        return bound.arguments

    return bind


class FakeRedis:
    """Every message a task publishes, and the keys it polls."""

    def __init__(self):
        self.messages = []  # (channel, message)
        self.keys = {}
        self.on_publish = None  # async hook(channel, message)

    def install(self, monkeypatch):
        import redis.asyncio as aioredis

        bind_message = _binding(job_support.publish_message)
        bind_progress = _binding(job_support.publish_progress)
        bind_sync = _binding(job_support.publish_sync)
        bind_flag = _binding(job_support.flag_is_set)
        bind_delete = _binding(job_support.delete_keys)
        bind_set = _binding(aioredis.Redis.set)
        bind_close = _binding(aioredis.Redis.close)
        fake = self

        async def publish_message(*args, **kwargs):
            a = bind_message(*args, **kwargs)
            message = json.loads(json.dumps(dict(a["message"])))
            fake.messages.append((a["channel"], message))
            if fake.on_publish is not None:
                await fake.on_publish(a["channel"], message)

        async def publish_progress(*args, **kwargs):
            a = bind_progress(*args, **kwargs)
            message = {
                "type": "progress",
                "progress": a["progress"],
                "stage": a["stage"],
                "status": a["status"],
            }
            if a["error"]:
                message["error"] = a["error"]
            await publish_message(a["channel"], message)

        def publish_sync(*args, **kwargs):
            a = bind_sync(*args, **kwargs)
            fake.messages.append((a["channel"], json.loads(json.dumps(a["message"]))))

        def flag_is_set(*args, **kwargs):
            a = bind_flag(*args, **kwargs)
            return bool(fake.keys.get(a["key"]))

        def delete_keys(*args, **kwargs):
            for key in bind_delete(*args, **kwargs)["keys"]:
                fake.keys.pop(key, None)

        class Client:
            async def set(self, *args, **kwargs):
                a = bind_set(self, *args, **kwargs)
                fake.keys[a["name"]] = a["value"]
                return True

            async def close(self, *args, **kwargs):
                bind_close(self, *args, **kwargs)

        monkeypatch.setattr(job_support, "publish_message", publish_message)
        monkeypatch.setattr(job_support, "publish_progress", publish_progress)
        monkeypatch.setattr(job_support, "publish_sync", publish_sync)
        monkeypatch.setattr(job_support, "flag_is_set", flag_is_set)
        monkeypatch.setattr(job_support, "delete_keys", delete_keys)
        monkeypatch.setattr(aioredis, "from_url", lambda *a, **k: Client())
        return self

    def on(self, channel):
        return [m for c, m in self.messages if c == channel]


class FakeMinio:
    """An object store in a dict, behind MinIOStorageService's own signatures.

    Each method answers the way the real one does at its edges: a missing
    object makes ``download_file`` return False and ``get_file_content`` raise
    FileNotFoundError; ``delete_file`` answers a bool rather than raising.
    """

    def __init__(self, delete_answers=None, fail_uploads=None):
        self.objects = {}
        self.content_types = {}
        self.deleted = []
        self.delete_answers = dict(delete_answers or {})
        self.fail_uploads = fail_uploads

    def install(self, monkeypatch):
        cls = MinIOStorageService
        b = {
            name: _binding(getattr(cls, name))
            for name in (
                "initialize",
                "upload_file",
                "upload_file_from_path",
                "upload_to_path",
                "delete_file",
                "download_file",
                "get_file_content",
            )
        }
        store = self

        async def initialize(*args, **kwargs):
            b["initialize"](*args, **kwargs)

        async def upload_file(*args, **kwargs):
            a = b["upload_file"](*args, **kwargs)
            assert isinstance(a["content"], (bytes, bytearray))
            path = a["self"]._get_object_path(a["document_id"], a["filename"])
            store._put(path, a["content"], a["content_type"])
            return path

        async def upload_file_from_path(*args, **kwargs):
            a = b["upload_file_from_path"](*args, **kwargs)
            path = a["self"]._get_object_path(a["document_id"], a["filename"])
            with open(a["file_path"], "rb") as handle:
                store._put(path, handle.read(), a["content_type"])
            return path

        async def upload_to_path(*args, **kwargs):
            a = b["upload_to_path"](*args, **kwargs)
            assert isinstance(a["content"], (bytes, bytearray))
            assert isinstance(a["object_path"], str)
            store._put(a["object_path"], a["content"], a["content_type"])
            return a["object_path"]

        async def delete_file(*args, **kwargs):
            a = b["delete_file"](*args, **kwargs)
            path = a["object_path"]
            store.deleted.append(path)
            answer = store.delete_answers.get(path, True)
            if answer:
                store.objects.pop(path, None)
            return answer

        async def download_file(*args, **kwargs):
            a = b["download_file"](*args, **kwargs)
            if a["object_path"] not in store.objects:
                return False
            with open(a["local_path"], "wb") as handle:
                handle.write(store.objects[a["object_path"]])
            return True

        async def get_file_content(*args, **kwargs):
            a = b["get_file_content"](*args, **kwargs)
            if a["object_path"] not in store.objects:
                raise FileNotFoundError(f"File not found: {a['object_path']}")
            return store.objects[a["object_path"]]

        for name, fn in (
            ("initialize", initialize),
            ("upload_file", upload_file),
            ("upload_file_from_path", upload_file_from_path),
            ("upload_to_path", upload_to_path),
            ("delete_file", delete_file),
            ("download_file", download_file),
            ("get_file_content", get_file_content),
        ):
            monkeypatch.setattr(cls, name, fn)

        # A method patched on the shared `storage_service` *instance* by an
        # earlier test (monkeypatch.setattr) is restored as an instance
        # attribute, which then shadows any class-level patch for the rest of
        # the run -- this fake was bypassed and the real MinIO client called.
        # Lift those leftovers for this test; monkeypatch puts them back.
        from app.services.storage_service import storage_service as _singleton

        for _name in list(vars(_singleton)):
            if callable(getattr(cls, _name, None)) and _name not in ("_get_client",):
                monkeypatch.delattr(_singleton, _name)
        return self

    def _put(self, path, content, content_type):
        if self.fail_uploads:
            raise self.fail_uploads
        self.objects[path] = bytes(content)
        self.content_types[path] = content_type


class FakeModel:
    """``LLMService.generate_response``, answered by ``answer(prompt)``."""

    def __init__(self, answer):
        self.answer = answer
        self.calls = []

    def install(self, monkeypatch):
        bind = _binding(LLMService.generate_response)
        model = self

        async def generate_response(*args, **kwargs):
            a = bind(*args, **kwargs)
            prompt = a["query"] or a["prompt"] or a["user_message"]
            assert isinstance(prompt, str) and prompt
            model.calls.append(a)
            reply = model.answer(prompt)
            if isinstance(reply, BaseException):
                raise reply
            return reply

        monkeypatch.setattr(LLMService, "generate_response", generate_response)
        return self

    def prompts(self):
        return [c["query"] or c["prompt"] or c["user_message"] for c in self.calls]


class FakeVectorStore:
    """Chunks in a list. ``search`` honours ``limit`` and ``document_ids``."""

    def __init__(self, chunks):
        self.chunks = chunks  # ranked best-first
        self.searches = []

    def install(self, monkeypatch):
        bind_init = _binding(VectorStoreService.initialize)
        bind_search = _binding(VectorStoreService.search)
        store = self

        async def initialize(*args, **kwargs):
            bind_init(*args, **kwargs)

        async def search(*args, **kwargs):
            a = bind_search(*args, **kwargs)
            store.searches.append(a)
            hits = store.chunks
            if a["document_ids"]:
                wanted = {str(d) for d in a["document_ids"]}
                hits = [h for h in hits if h["metadata"]["document_id"] in wanted]
            return [dict(h) for h in hits[: a["limit"]]]

        monkeypatch.setattr(VectorStoreService, "initialize", initialize)
        monkeypatch.setattr(VectorStoreService, "search", search)
        return self


def _png():
    from PIL import Image

    buf = io.BytesIO()
    Image.new("RGB", (8, 8), "white").save(buf, "PNG")
    return buf.getvalue()


README = "# Widget\n\nInstall with make install.\n\n- Fast stride prefetching\n"


class FakeGitHub:
    """The GitHub REST API for one repository, plus a diagram renderer."""

    OWNER, REPO = "acme", "widget"

    def __init__(self):
        self.requests = []
        self.png = _png()

    def install(self, monkeypatch):
        real = httpx.AsyncClient
        transport = httpx.MockTransport(self.handle)

        class Client(real):
            def __init__(self, *args, **kwargs):
                kwargs["transport"] = transport
                super().__init__(*args, **kwargs)

        monkeypatch.setattr(httpx, "AsyncClient", Client)
        return self

    def hosts(self):
        return [r.url.host for r in self.requests]

    def handle(self, request):
        self.requests.append(request)
        if request.url.host != "api.github.com":
            return httpx.Response(200, content=self.png)  # a diagram renderer
        path = request.url.path
        base = f"/repos/{self.OWNER}/{self.REPO}"
        if path == "/user":
            if request.headers.get("authorization"):
                return httpx.Response(200, json={"login": "me"})
            return httpx.Response(401, json={"message": "Requires authentication"})
        if not path.startswith(base):
            return httpx.Response(404, json={"message": "Not Found"})
        rest = path[len(base) :]
        if rest == "":
            return httpx.Response(
                200,
                json={
                    "name": self.REPO,
                    "full_name": f"{self.OWNER}/{self.REPO}",
                    "description": "A cache prefetching library",
                    "html_url": f"https://github.com/{self.OWNER}/{self.REPO}",
                    "default_branch": "main",
                    "stargazers_count": 120,
                    "forks_count": 7,
                    "watchers_count": 9,
                    "license": {"name": "MIT License"},
                    "language": "C",
                    "created_at": "2024-01-02T03:04:05Z",
                    "updated_at": "2026-09-30T00:00:00Z",
                },
            )
        if rest == "/readme":
            if "html" in request.headers.get("accept", ""):
                return httpx.Response(200, text="<h1>Widget</h1>")
            return httpx.Response(
                200,
                json={
                    "name": "README.md",
                    "encoding": "base64",
                    "content": base64.b64encode(README.encode()).decode(),
                },
            )
        if rest == "/git/trees/HEAD":
            return httpx.Response(
                200,
                json={
                    "tree": [
                        {"path": "src", "type": "tree"},
                        {"path": "src/stride.c", "type": "blob", "size": 900},
                        {"path": "docs", "type": "tree"},
                        {"path": "README.md", "type": "blob", "size": 60},
                    ]
                },
            )
        if rest == "/commits":
            return httpx.Response(
                200,
                json=[
                    {
                        "sha": "abcdef1234567890",
                        "html_url": "https://github.com/acme/widget/commit/abcdef1",
                        "commit": {
                            "message": "Fix stride prefetch distance\n\nLonger body",
                            "author": {
                                "name": "Ada",
                                "email": "ada@example.com",
                                "date": "2026-09-01T10:00:00Z",
                            },
                        },
                    }
                ],
            )
        if rest == "/issues":
            return httpx.Response(
                200,
                json=[
                    {
                        "number": 7,
                        "title": "ISB never engages",
                        "state": "open",
                        "user": {"login": "bob"},
                        "created_at": "2026-09-02T00:00:00Z",
                        "labels": [{"name": "bug"}],
                    },
                    {
                        "number": 9,
                        "title": "Add hybrid prefetcher",
                        "pull_request": {"url": "x"},
                        "user": {"login": "carol"},
                    },
                ],
            )
        if rest == "/pulls":
            return httpx.Response(
                200,
                json=[
                    {
                        "number": 9,
                        "title": "Add hybrid prefetcher",
                        "state": "open",
                        "user": {"login": "carol"},
                        "created_at": "2026-09-03T00:00:00Z",
                        "labels": [],
                        "head": {"ref": "hybrid"},
                        "base": {"ref": "main"},
                    }
                ],
            )
        if rest == "/languages":
            return httpx.Response(200, json={"C": 9000, "Python": 1000})
        if rest == "/contributors":
            return httpx.Response(200, json=[{"login": "ada", "contributions": 42}])
        return httpx.Response(404, json={"message": "Not Found"})


ARCHITECTURE = (
    "The project is split into a core prefetch library and command line tools. "
    "Each component reports through a shared statistics module."
)


def _insights_answer(prompt):
    if "architecture summary" in prompt:
        return ARCHITECTURE
    if "key features" in prompt:
        return "- Stride prefetching\n- Hybrid stream detection"
    if "technology stack" in prompt:
        return "- C\n- CMake"
    return "unexpected prompt"


@pytest.fixture
def redis_fake(monkeypatch):
    return FakeRedis().install(monkeypatch)


@pytest.fixture
def minio(monkeypatch):
    return FakeMinio().install(monkeypatch)


@pytest.fixture
def task_sessions(db_session, monkeypatch):
    factory = async_sessionmaker(
        db_session.bind, class_=AsyncSession, expire_on_commit=False
    )
    for module in TASK_MODULES:
        monkeypatch.setattr(module, "create_celery_session", lambda: factory)
    return factory


@pytest.fixture
def output_dir(tmp_path, monkeypatch):
    out = tmp_path / "training_outputs"
    monkeypatch.setattr(settings, "TRAINING_OUTPUT_DIR", str(out))
    return out


async def _reload(db, model, row_id):
    """The row as the task left it."""
    return (
        await db.execute(
            select(model)
            .where(model.id == row_id)
            .execution_options(populate_existing=True)
        )
    ).scalar_one_or_none()


def _docx_text(blob):
    import docx

    document = docx.Document(io.BytesIO(blob))
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


def _pptx(blob):
    import pptx

    return pptx.Presentation(io.BytesIO(blob))


def _pptx_text(blob):
    out = []
    for slide in _pptx(blob).slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                out.append(shape.text_frame.text)
    return "\n".join(out)


# --------------------------------------------------------------------------
# Rows
# --------------------------------------------------------------------------


async def _source(db, **overrides):
    fields = {"name": f"src-{uuid4().hex[:8]}", "source_type": "file", "config": {}}
    fields.update(overrides)
    source = DocumentSource(**fields)
    db.add(source)
    await db.commit()
    await db.refresh(source)
    return source


async def _document(db, **overrides):
    source = await _source(db)
    fields = {
        "title": "Widget Spec",
        "content": "The widget prefetches cache lines ahead of a stride.",
        "content_hash": uuid4().hex,
        "source_id": source.id,
        "source_identifier": "spec.pdf",
        "file_type": "pdf",
    }
    fields.update(overrides)
    doc = Document(**fields)
    db.add(doc)
    await db.commit()
    await db.refresh(doc)
    return doc


SAMPLES = [
    {
        "instruction": f"What does prefetcher {i} do?",
        "input": "",
        "output": f"Prefetcher {i} predicts the next miss from a stride.",
    }
    for i in range(12)
]


async def _dataset(db, user, *, samples=SAMPLES, file_path="auto", **overrides):
    fields = {
        "name": "Prefetch QA",
        "user_id": user.id,
        "status": "ready",
        "sample_count": len(samples),
    }
    fields.update(overrides)
    dataset = TrainingDataset(**fields)
    db.add(dataset)
    await db.flush()
    for i, content in enumerate(samples):
        db.add(
            DatasetSample(
                dataset_id=dataset.id,
                sample_index=i,
                content=content,
                input_tokens=len(content["instruction"]) // 4,
                output_tokens=len(content["output"]) // 4,
            )
        )
    if file_path == "auto":
        file_path = f"training/datasets/{dataset.id}/dataset.jsonl"
    dataset.file_path = file_path
    await db.commit()
    await db.refresh(dataset)
    return dataset


def _jsonl(samples=SAMPLES):
    return "\n".join(json.dumps(s) for s in samples).encode()


async def _training_job(db, user, dataset, **overrides):
    fields = {
        "name": "Prefetch Adapter",
        "training_method": "lora",
        "training_backend": "simulated",
        "base_model": "llama3.2:1b",
        "dataset_id": dataset.id,
        "hyperparameters": {"num_epochs": 1, "batch_size": 4},
        "user_id": user.id,
        "status": "queued",
        "progress": 0,
    }
    fields.update(overrides)
    job = TrainingJob(**fields)
    db.add(job)
    await db.commit()
    await db.refresh(job)
    return job


def _template_docx(sections, *, title=None):
    """A DOCX whose headings are ``sections`` (title -> placeholder or None)."""
    import docx

    document = docx.Document()
    if title:
        document.add_heading(title, level=0)
    for heading, placeholder in sections.items():
        document.add_heading(heading, level=1)
        if placeholder is not None:
            document.add_paragraph(placeholder)
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


SPEC_TEMPLATE = {
    "Overview": "Describe the product here.",
    "Thermal Limits": "State the operating temperature range.",
}


async def _template_job(db, user, docs, store, template=None, **overrides):
    path = f"templates/{uuid4()}/spec.docx"
    if template is not None:
        store.objects[path] = template
    fields = {
        "user_id": user.id,
        "template_file_path": path,
        "template_filename": "spec.docx",
        "source_document_ids": [str(d.id) for d in docs],
        "status": "pending",
        "progress": 0,
    }
    fields.update(overrides)
    job = TemplateJob(**fields)
    db.add(job)
    await db.commit()
    await db.refresh(job)
    return job


def _chunk(doc, content, title=None):
    return {
        "id": uuid4().hex,
        "content": content,
        "metadata": {"document_id": str(doc.id), "title": title or doc.title},
        "score": 0.9,
    }


def _section_answer(prompt):
    if 'titled "Overview"' in prompt:
        return "The Widget is a hardware stride prefetcher."
    if 'titled "Thermal Limits"' in prompt:
        return "It operates from -40C to 85C."
    return "Generic section content."


ALL_SECTIONS = [
    "overview",
    "readme",
    "file_structure",
    "commits",
    "issues",
    "pull_requests",
    "code_stats",
    "contributors",
    "architecture",
    "technology_stack",
]


async def _repo_job(db, user, **overrides):
    fields = {
        "user_id": user.id,
        "adhoc_url": "https://github.com/acme/widget",
        "repo_name": "acme/widget",
        "repo_url": "https://github.com/acme/widget",
        "repo_type": "github",
        "output_format": "docx",
        "title": "Widget Report",
        "sections": list(ALL_SECTIONS),
        "include_diagrams": True,
        "style": "professional",
        "status": "pending",
        "progress": 0,
    }
    fields.update(overrides)
    job = RepoReportJob(**fields)
    db.add(job)
    await db.commit()
    await db.refresh(job)
    return job


# --------------------------------------------------------------------------
# The synchronous Celery functions, each on its own engine
# --------------------------------------------------------------------------


@pytest.fixture
def file_db(tmp_path, monkeypatch):
    """A file database reachable from any event loop, as Postgres is."""
    url = f"sqlite+aiosqlite:///{tmp_path / 'tasks.db'}"

    def factory():
        engine = create_async_engine(url, poolclass=NullPool)
        return async_sessionmaker(engine, class_=AsyncSession, expire_on_commit=False)

    for module in TASK_MODULES:
        monkeypatch.setattr(module, "create_celery_session", factory)

    async def create():
        engine = create_async_engine(url, poolclass=NullPool)
        async with engine.begin() as conn:
            await conn.run_sync(Base.metadata.create_all)
        await engine.dispose()

    asyncio.run(create())

    def run(coro_fn):
        async def go():
            async with factory()() as db:
                return await coro_fn(db)

        return asyncio.run(go())

    return run


def _seed_user(run):
    async def make(db):
        return await AuthService().create_user(
            username="syncuser",
            email="sync@example.com",
            password="testpassword123",
            full_name="Sync User",
            db=db,
        )

    return run(make)


def _get(run, model, row_id):
    async def fetch(db):
        return (
            await db.execute(select(model).where(model.id == row_id))
        ).scalar_one_or_none()

    return run(fetch)


def _all(run, model, *where):
    async def fetch(db):
        return list((await db.execute(select(model).where(*where))).scalars().all())

    return run(fetch)


# ==========================================================================
# Training
# ==========================================================================


async def _ready_training(db, user, store, **job_overrides):
    dataset = await _dataset(db, user)
    store.objects[dataset.file_path] = _jsonl()
    job = await _training_job(db, user, dataset, **job_overrides)
    return dataset, job


async def test_training_job_completes_and_registers_its_adapter(
    db_session, test_user, task_sessions, redis_fake, minio, output_dir
):
    _, job = await _ready_training(db_session, test_user, minio)

    await training_tasks._execute_training_async(str(job.id), str(test_user.id))

    row = await _reload(db_session, TrainingJob, job.id)
    assert row.status == "completed", row.error
    assert row.progress == 100
    assert row.error is None
    assert row.started_at is not None and row.completed_at is not None
    # 12 samples, batch 4, 1 epoch -> the simulated trainer's floor of 10 steps.
    assert row.current_step == row.total_steps == 10
    assert row.final_metrics["trainer"] == "simulated"
    assert row.final_metrics["samples"] == 12  # it trained on the downloaded file

    prefix = f"training/adapters/{test_user.id}/{job.id}"
    assert sorted(minio.objects) == sorted(
        [
            next(p for p in minio.objects if p.startswith("training/datasets/")),
            f"{prefix}/adapter.json",
            f"{prefix}/README.txt",
        ]
    )
    artifact = json.loads(minio.objects[f"{prefix}/adapter.json"])
    assert artifact["job_id"] == str(job.id)
    assert artifact["base_model"] == "llama3.2:1b"

    adapter = await _reload(db_session, ModelAdapter, row.output_adapter_id)
    assert adapter.adapter_path == prefix
    assert adapter.training_job_id == job.id
    assert adapter.user_id == test_user.id
    assert adapter.status == AdapterStatus.READY.value
    assert adapter.adapter_size == len(minio.objects[f"{prefix}/adapter.json"]) + len(
        minio.objects[f"{prefix}/README.txt"]
    )

    channel = redis_fake.on(f"training_job:{job.id}:progress")
    statuses = [m["status"] for m in channel]
    assert statuses[0] == "preparing"
    assert statuses[-2:] == ["saving", "completed"]
    assert channel[-1]["progress"] == 100
    steps = [m["current_step"] for m in channel if "current_step" in m]
    assert steps == list(range(1, 11))


async def test_training_metrics_record_every_reported_step(
    db_session, test_user, task_sessions, redis_fake, minio, output_dir
):
    _, job = await _ready_training(db_session, test_user, minio)

    await training_tasks._execute_training_async(str(job.id), str(test_user.id))

    row = await _reload(db_session, TrainingJob, job.id)
    assert row.status == "completed", row.error
    reported = [
        m["current_loss"]
        for m in redis_fake.on(f"training_job:{job.id}:progress")
        if "current_loss" in m
    ]
    assert len(reported) == 10
    assert row.training_metrics["loss_history"] == reported
    assert row.training_metrics["current_loss"] == reported[-1]
    assert row.training_metrics["best_loss"] == min(reported)


@pytest.mark.parametrize(
    "overrides,store_dataset,cause",
    [
        ({"training_backend": "modal"}, True, "Trainer 'modal' not available"),
        ({}, False, "File not found: training/datasets/"),
        (
            {"hyperparameters": {"num_epochs": "three"}},
            True,
            "invalid literal for int() with base 10: 'three'",
        ),
    ],
    ids=["unavailable-backend", "dataset-file-missing", "trainer-raises"],
)
async def test_training_job_that_cannot_run_fails_naming_the_cause(
    db_session,
    test_user,
    task_sessions,
    redis_fake,
    minio,
    output_dir,
    overrides,
    store_dataset,
    cause,
):
    dataset = await _dataset(db_session, test_user)
    if store_dataset:
        minio.objects[dataset.file_path] = _jsonl()
    job = await _training_job(db_session, test_user, dataset, **overrides)

    await training_tasks._execute_training_async(str(job.id), str(test_user.id))

    row = await _reload(db_session, TrainingJob, job.id)
    assert row.status == "failed"
    assert row.error.startswith(cause)
    assert row.completed_at is not None
    assert row.output_adapter_id is None
    assert not any(p.startswith("training/adapters/") for p in minio.objects)
    last = redis_fake.on(f"training_job:{job.id}:progress")[-1]
    assert last["status"] == "failed" and last["error"] == row.error


async def test_failed_training_leaves_no_dataset_copy_behind(
    db_session, test_user, task_sessions, redis_fake, minio, output_dir, tmp_path
):
    scratch = tmp_path / "scratch"
    scratch.mkdir()
    import tempfile

    original = tempfile.tempdir
    tempfile.tempdir = str(scratch)
    try:
        _, job = await _ready_training(
            db_session, test_user, minio, hyperparameters={"num_epochs": "three"}
        )
        await training_tasks._execute_training_async(str(job.id), str(test_user.id))
    finally:
        tempfile.tempdir = original

    row = await _reload(db_session, TrainingJob, job.id)
    assert row.status == "failed"
    assert os.listdir(scratch) == []


async def test_training_job_exports_an_unexported_dataset_first(
    db_session, test_user, task_sessions, redis_fake, minio, output_dir
):
    dataset = await _dataset(db_session, test_user, file_path=None)
    job = await _training_job(db_session, test_user, dataset)

    await training_tasks._execute_training_async(str(job.id), str(test_user.id))

    row = await _reload(db_session, TrainingJob, job.id)
    assert row.status == "completed", row.error
    exported = await _reload(db_session, TrainingDataset, dataset.id)
    assert exported.file_path == f"training/datasets/{dataset.id}/dataset.jsonl"
    assert minio.objects[exported.file_path] == _jsonl()
    assert exported.file_size == len(_jsonl())


async def test_training_job_cancelled_mid_run_stays_cancelled(
    db_session, test_user, task_sessions, redis_fake, minio, output_dir
):
    _, job = await _ready_training(db_session, test_user, minio)
    cancelled = []

    async def cancel_at_step_three(channel, message):
        if message.get("current_step") == 3 and not cancelled:
            async with task_sessions() as other:
                await training_service.cancel_job(other, job.id, test_user.id)
            cancelled.append(True)

    redis_fake.on_publish = cancel_at_step_three

    await training_tasks._execute_training_async(str(job.id), str(test_user.id))

    assert cancelled, "the run never reached step 3"
    assert redis_fake.keys.get(f"training_job:{job.id}:cancel") == "1"
    row = await _reload(db_session, TrainingJob, job.id)
    assert row.output_adapter_id is None
    assert row.status == "cancelled"


async def test_missing_training_job_is_a_no_op(
    db_session, test_user, task_sessions, redis_fake, minio, output_dir
):
    await training_tasks._execute_training_async(str(uuid4()), str(test_user.id))

    assert redis_fake.messages == []
    assert minio.objects == {}


async def test_cancelled_training_job_is_left_alone(
    db_session, test_user, task_sessions, redis_fake, minio, output_dir
):
    _, job = await _ready_training(db_session, test_user, minio, status="cancelled")

    await training_tasks._execute_training_async(str(job.id), str(test_user.id))

    row = await _reload(db_session, TrainingJob, job.id)
    assert row.status == "cancelled" and row.started_at is None
    assert redis_fake.messages == []


def test_execute_training_job_task_run_completes_a_job(
    file_db, redis_fake, minio, output_dir
):
    user = _seed_user(file_db)

    async def seed(db):
        return await _ready_training(db, user, minio)

    _, job = file_db(seed)

    training_tasks.execute_training_job_task.run(str(job.id), str(user.id))

    row = _get(file_db, TrainingJob, job.id)
    assert row.status == "completed", row.error
    assert row.output_adapter_id is not None
    assert f"training/adapters/{user.id}/{job.id}/adapter.json" in minio.objects


def test_execute_training_job_task_run_on_a_missing_job(file_db, redis_fake, minio):
    training_tasks.execute_training_job_task.run(str(uuid4()), str(uuid4()))
    assert redis_fake.messages == []


def test_export_dataset_task_writes_the_jsonl(file_db, minio):
    user = _seed_user(file_db)
    dataset = file_db(lambda db: _dataset(db, user, file_path=None))

    path = training_tasks.export_dataset_task.run(str(dataset.id), str(user.id))

    assert path == f"training/datasets/{dataset.id}/dataset.jsonl"
    assert minio.objects[path] == _jsonl()


# ---- validate_dataset_task ------------------------------------------------


def test_validate_dataset_task_marks_a_good_dataset_ready(file_db):
    user = _seed_user(file_db)
    dataset = file_db(lambda db: _dataset(db, user, status="draft"))

    assert training_tasks.validate_dataset_task.run(str(dataset.id), str(user.id))

    row = _get(file_db, TrainingDataset, dataset.id)
    assert row.status == "ready"
    assert row.is_validated is True
    assert row.validation_errors is None


def test_validate_dataset_task_refuses_a_dataset_too_small_to_train(file_db):
    user = _seed_user(file_db)
    dataset = file_db(lambda db: _dataset(db, user, samples=SAMPLES[:3]))

    valid = training_tasks.validate_dataset_task.run(str(dataset.id), str(user.id))

    assert valid is False
    row = _get(file_db, TrainingDataset, dataset.id)
    assert row.status == "error"
    assert row.is_validated is False
    assert [e["code"] for e in row.validation_errors] == ["MIN_SAMPLES"]
    assert "only 3 samples" in row.validation_errors[0]["message"]


def test_validate_dataset_task_on_a_missing_dataset_names_it(file_db):
    missing = uuid4()
    with pytest.raises(ValueError, match=f"Dataset {missing} not found"):
        training_tasks.validate_dataset_task.run(str(missing), str(uuid4()))


# ---- generate_dataset_from_documents_task ---------------------------------


QA_REPLY = json.dumps(
    [
        {"question": "What does the widget prefetch?", "answer": "Cache lines."},
        {"instruction": "Name its predictor.", "output": "A stride table."},
    ]
)


def test_generate_dataset_task_builds_samples_from_each_document(file_db, monkeypatch):
    model = FakeModel(lambda prompt: QA_REPLY).install(monkeypatch)
    user = _seed_user(file_db)
    first = file_db(lambda db: _document(db))
    second = file_db(
        lambda db: _document(
            db, title="Tiling", content=None, summary="Tiling bounds the working set."
        )
    )
    empty = file_db(lambda db: _document(db, title="Empty", content=None))

    dataset_id = training_tasks.generate_dataset_from_documents_task.run(
        str(user.id),
        "Widget QA",
        "From the spec",
        [str(first.id), str(second.id), str(empty.id)],
        samples_per_document=2,
    )

    row = _get(file_db, TrainingDataset, UUID(dataset_id))
    assert row.name == "Widget QA" and row.user_id == user.id
    assert row.status == "draft"
    assert row.sample_count == 4
    assert sorted(row.source_document_ids) == sorted(
        [str(first.id), str(second.id), str(empty.id)]
    )
    samples = _all(file_db, DatasetSample, DatasetSample.dataset_id == row.id)
    assert sorted(s.sample_index for s in samples) == [0, 1, 2, 3]
    assert {s.source_document_id for s in samples} == {first.id, second.id}
    assert {s.content["instruction"] for s in samples} == {
        "What does the widget prefetch?",
        "Name its predictor.",
    }
    assert row.token_count == sum(s.input_tokens + s.output_tokens for s in samples)
    # One call per document with text; the empty one costs nothing.
    assert len(model.calls) == 2
    prompts = model.prompts()
    assert any("prefetches cache lines ahead of a stride" in p for p in prompts)
    assert any("Tiling bounds the working set." in p for p in prompts)
    assert all("Generate exactly 2 training samples" in p for p in prompts)


def test_generate_dataset_task_keeps_what_one_bad_reply_did_not_spoil(
    file_db, monkeypatch
):
    FakeModel(
        lambda prompt: "Sorry, I cannot." if "Tiling" in prompt else QA_REPLY
    ).install(monkeypatch)
    user = _seed_user(file_db)
    good = file_db(lambda db: _document(db))
    bad = file_db(lambda db: _document(db, content="Tiling bounds the working set."))

    dataset_id = training_tasks.generate_dataset_from_documents_task.run(
        str(user.id), "Widget QA", "", [str(good.id), str(bad.id)]
    )

    samples = _all(file_db, DatasetSample, DatasetSample.dataset_id == UUID(dataset_id))
    assert len(samples) == 2
    assert {s.source_document_id for s in samples} == {good.id}


# ---- cleanup_old_checkpoints_task -----------------------------------------


def _seed_checkpoints(file_db, minio, steps=(100, 200, 300, 400, 500)):
    user = _seed_user(file_db)

    async def seed(db):
        dataset = await _dataset(db, user)
        job = await _training_job(db, user, dataset, status="completed")
        for step in steps:
            path = f"training/checkpoints/{job.id}/step-{step}"
            minio.objects[path] = b"weights"
            db.add(
                TrainingCheckpoint(
                    job_id=job.id, step=step, checkpoint_path=path, loss=1.0
                )
            )
        await db.commit()
        return job

    return file_db(seed)


def test_cleanup_old_checkpoints_keeps_the_latest(file_db, monkeypatch):
    store = FakeMinio().install(monkeypatch)
    job = _seed_checkpoints(file_db, store)

    deleted = training_tasks.cleanup_old_checkpoints_task.run(str(job.id), keep_last=2)

    assert deleted == 3
    left = _all(file_db, TrainingCheckpoint, TrainingCheckpoint.job_id == job.id)
    assert sorted(c.step for c in left) == [400, 500]
    assert sorted(store.deleted) == [
        f"training/checkpoints/{job.id}/step-{s}" for s in (100, 200, 300)
    ]
    assert sorted(store.objects) == [
        f"training/checkpoints/{job.id}/step-{s}" for s in (400, 500)
    ]


def test_cleanup_old_checkpoints_of_an_unknown_job_deletes_nothing(
    file_db, monkeypatch
):
    store = FakeMinio().install(monkeypatch)
    assert training_tasks.cleanup_old_checkpoints_task.run(str(uuid4())) == 0
    assert store.deleted == []


def test_cleanup_keeps_a_checkpoint_whose_file_survived(file_db, monkeypatch):
    store = FakeMinio().install(monkeypatch)
    job = _seed_checkpoints(file_db, store, steps=(100, 200))
    stuck = f"training/checkpoints/{job.id}/step-100"
    store.delete_answers[stuck] = False

    training_tasks.cleanup_old_checkpoints_task.run(str(job.id), keep_last=1)

    assert store.deleted == [stuck]
    left = _all(file_db, TrainingCheckpoint, TrainingCheckpoint.job_id == job.id)
    assert sorted(c.step for c in left) == [100, 200]


# ==========================================================================
# Template fill
# ==========================================================================


async def _run_fill(job):
    return await template_tasks._async_fill_template(None, str(job.id))


async def test_template_fill_writes_each_section_from_its_sources(
    db_session, test_user, task_sessions, redis_fake, minio, monkeypatch
):
    doc = await _document(db_session)
    vectors = FakeVectorStore(
        [
            _chunk(doc, "The widget is a stride prefetcher for L2."),
            _chunk(doc, "Rated from -40C to 85C at full load."),
        ]
    ).install(monkeypatch)
    model = FakeModel(_section_answer).install(monkeypatch)
    job = await _template_job(
        db_session, test_user, [doc], minio, template=_template_docx(SPEC_TEMPLATE)
    )

    outcome = await _run_fill(job)

    row = await _reload(db_session, TemplateJob, job.id)
    assert row.status == "completed", row.error_message
    assert outcome == {
        "success": True,
        "job_id": str(job.id),
        "filled_filename": "filled_spec.docx",
        "filled_file_path": f"{job.id}/filled_spec.docx",
    }
    assert row.filled_file_path == f"{job.id}/filled_spec.docx"
    assert row.filled_filename == "filled_spec.docx"
    assert row.progress == 100 and row.completed_at is not None
    assert row.current_section is None
    assert [s["title"] for s in row.sections] == ["Overview", "Thermal Limits"]
    assert row.sections[0]["placeholder_text"] == "Describe the product here."

    assert minio.content_types[row.filled_file_path] == DOCX_MIME
    text = _docx_text(minio.objects[row.filled_file_path])
    assert "Overview" in text and "Thermal Limits" in text
    assert "The Widget is a hardware stride prefetcher." in text
    assert "It operates from -40C to 85C." in text
    assert "Describe the product here." not in text

    assert [s["query"] for s in vectors.searches] == ["Overview", "Thermal Limits"]
    assert len(model.calls) == 2
    assert all("Rated from -40C to 85C" in p for p in model.prompts())
    channel = redis_fake.on(f"template_progress:{job.id}")
    assert channel[-1] == {
        "type": "complete",
        "job_id": str(job.id),
        "result": {
            "filled_filename": "filled_spec.docx",
            "filled_file_path": row.filled_file_path,
        },
    }


async def test_template_fill_falls_back_to_document_text_without_chunks(
    db_session, test_user, task_sessions, redis_fake, minio, monkeypatch
):
    doc = await _document(db_session, content="Operating range: -40C to 85C.")
    FakeVectorStore([]).install(monkeypatch)
    model = FakeModel(_section_answer).install(monkeypatch)
    job = await _template_job(
        db_session,
        test_user,
        [doc],
        minio,
        template=_template_docx({"Thermal Limits": "Range here."}),
    )

    await _run_fill(job)

    row = await _reload(db_session, TemplateJob, job.id)
    assert row.status == "completed", row.error_message
    assert "From 'Widget Spec':\nOperating range: -40C to 85C." in model.prompts()[0]


async def test_template_without_sections_fails_naming_it(
    db_session, test_user, task_sessions, redis_fake, minio, monkeypatch
):
    import docx

    plain = docx.Document()
    plain.add_paragraph("Just a sentence with no heading at all, nothing more.")
    buf = io.BytesIO()
    plain.save(buf)
    doc = await _document(db_session)
    FakeVectorStore([]).install(monkeypatch)
    model = FakeModel(_section_answer).install(monkeypatch)
    job = await _template_job(
        db_session, test_user, [doc], minio, template=buf.getvalue()
    )

    outcome = await _run_fill(job)

    row = await _reload(db_session, TemplateJob, job.id)
    assert row.status == "failed"
    assert row.error_message == "No sections detected in template"
    assert outcome == {
        "success": False,
        "job_id": str(job.id),
        "error": "No sections detected in template",
    }
    assert model.calls == []
    assert not any(p.endswith("filled_spec.docx") for p in minio.objects)
    assert redis_fake.on(f"template_progress:{job.id}")[-1] == {
        "type": "error",
        "job_id": str(job.id),
        "error": "No sections detected in template",
    }


async def test_template_missing_from_storage_fails_naming_the_template(
    db_session, test_user, task_sessions, redis_fake, minio, monkeypatch
):
    doc = await _document(db_session)
    FakeVectorStore([]).install(monkeypatch)
    FakeModel(_section_answer).install(monkeypatch)
    job = await _template_job(db_session, test_user, [doc], minio, template=None)

    await _run_fill(job)

    row = await _reload(db_session, TemplateJob, job.id)
    assert row.status == "failed"
    assert job.template_file_path in (row.error_message or "")


async def test_template_fill_whose_model_fails_does_not_report_success(
    db_session, test_user, task_sessions, redis_fake, minio, monkeypatch
):
    doc = await _document(db_session)
    FakeVectorStore([_chunk(doc, "Rated to 85C.")]).install(monkeypatch)
    FakeModel(lambda prompt: RuntimeError("deepseek: 401 invalid api key")).install(
        monkeypatch
    )
    job = await _template_job(
        db_session, test_user, [doc], minio, template=_template_docx(SPEC_TEMPLATE)
    )

    await _run_fill(job)

    row = await _reload(db_session, TemplateJob, job.id)
    assert row.status == "failed"
    assert "deepseek: 401 invalid api key" in (row.error_message or "")


async def test_template_of_bare_headings_is_filled(
    db_session, test_user, task_sessions, redis_fake, minio, monkeypatch
):
    doc = await _document(db_session)
    FakeVectorStore([_chunk(doc, "Rated from -40C to 85C.")]).install(monkeypatch)
    model = FakeModel(_section_answer).install(monkeypatch)
    job = await _template_job(
        db_session,
        test_user,
        [doc],
        minio,
        template=_template_docx({"Overview": None, "Thermal Limits": None}),
    )

    await _run_fill(job)

    row = await _reload(db_session, TemplateJob, job.id)
    assert row.status == "completed", row.error_message
    assert len(model.calls) == 2
    text = _docx_text(minio.objects[row.filled_file_path])
    assert "The Widget is a hardware stride prefetcher." in text
    assert "It operates from -40C to 85C." in text


async def test_template_context_comes_from_the_source_documents_chunks(
    db_session, test_user, task_sessions, redis_fake, minio, monkeypatch
):
    other = await _document(db_session, title="Unrelated Manual")
    doc = await _document(
        db_session, content="Preamble. " * 400 + "Rated to 85C at full load."
    )
    noise = [_chunk(other, f"Thermal paste brand {i}") for i in range(10)]
    FakeVectorStore(noise + [_chunk(doc, "Rated to 85C at full load.")]).install(
        monkeypatch
    )
    model = FakeModel(_section_answer).install(monkeypatch)
    job = await _template_job(
        db_session,
        test_user,
        [doc],
        minio,
        template=_template_docx({"Thermal Limits": "Range here."}),
    )

    await _run_fill(job)

    row = await _reload(db_session, TemplateJob, job.id)
    assert row.status == "completed", row.error_message
    assert "Rated to 85C at full load." in model.prompts()[0]


async def test_missing_template_job_reports_it(
    db_session, test_user, task_sessions, redis_fake, minio
):
    missing = str(uuid4())

    outcome = await template_tasks._async_fill_template(None, missing)

    error = f"Template job {missing} not found"
    assert outcome == {"success": False, "job_id": missing, "error": error}
    assert redis_fake.messages == [
        (
            f"template_progress:{missing}",
            {"type": "error", "job_id": missing, "error": error},
        )
    ]


def test_fill_template_task_run_completes_a_job(
    file_db, redis_fake, minio, monkeypatch
):
    FakeVectorStore([]).install(monkeypatch)
    FakeModel(_section_answer).install(monkeypatch)
    user = _seed_user(file_db)

    async def seed(db):
        doc = await _document(db)
        return await _template_job(
            db, user, [doc], minio, template=_template_docx(SPEC_TEMPLATE)
        )

    job = file_db(seed)

    outcome = template_tasks.fill_template.run(str(job.id))

    assert outcome["success"] is True
    row = _get(file_db, TemplateJob, job.id)
    assert row.status == "completed", row.error_message
    assert "It operates from -40C to 85C." in _docx_text(
        minio.objects[row.filled_file_path]
    )


# ==========================================================================
# Repository reports
# ==========================================================================


@pytest.fixture
def github(monkeypatch):
    return FakeGitHub().install(monkeypatch)


async def _run_report(job, user):
    await repo_report_tasks._generate_repo_report_async(str(job.id), str(user.id))


async def test_repo_report_docx_carries_every_requested_section(
    db_session, test_user, task_sessions, redis_fake, minio, github, monkeypatch
):
    model = FakeModel(_insights_answer).install(monkeypatch)
    job = await _repo_job(db_session, test_user)

    await _run_report(job, test_user)

    row = await _reload(db_session, RepoReportJob, job.id)
    assert row.status == "completed", row.error
    assert row.file_path == f"repo_reports/{test_user.id}/acme_widget_{job.id}.docx"
    blob = minio.objects[row.file_path]
    assert row.file_size == len(blob)
    assert minio.content_types[row.file_path] == DOCX_MIME
    assert row.progress == 100 and row.current_stage == "Completed"
    assert row.started_at is not None and row.completed_at is not None
    assert row.analysis_data["repo_info"]["full_name"] == "acme/widget"
    assert row.analysis_data["repo_info"]["stars"] == 120

    text = _docx_text(blob)
    for expected in (
        "Widget Report",
        "A cache prefetching library",
        "MIT License",
        "Install with make install.",
        "src/",
        "Fix stride prefetch distance",
        "ISB never engages",
        "Add hybrid prefetcher",
        "hybrid → main",
        "90.0%",
        "ada",
        ARCHITECTURE,
        "Hybrid stream detection",
        "CMake",
    ):
        assert expected in text, expected
    assert "#9" not in text.split("Open Pull Requests")[0]  # a PR is not an issue
    assert len(model.calls) == 3
    assert "File structure:\nwidget/" in model.prompts()[0]
    final = redis_fake.on(f"repo_report:{job.id}:progress")[-1]
    assert final["status"] == "completed" and final["progress"] == 100


async def test_repo_report_from_a_source_uses_its_token_and_builds_a_pdf(
    db_session, test_user, task_sessions, redis_fake, minio, github, monkeypatch
):
    FakeModel(_insights_answer).install(monkeypatch)
    source = await _source(
        db_session,
        source_type="github",
        config={"repos": ["acme/widget"], "token": "ghp_secret"},
    )
    job = await _repo_job(
        db_session,
        test_user,
        source_id=source.id,
        adhoc_url=None,
        output_format="pdf",
    )

    await _run_report(job, test_user)

    row = await _reload(db_session, RepoReportJob, job.id)
    assert row.status == "completed", row.error
    assert row.file_path.endswith(".pdf")
    assert minio.content_types[row.file_path] == "application/pdf"
    assert minio.objects[row.file_path].startswith(b"%PDF")
    auth = {r.headers.get("authorization") for r in github.requests}
    assert auth == {"Bearer ghp_secret"}


async def test_repo_presentation_has_its_slides_and_diagram(
    db_session, test_user, task_sessions, redis_fake, minio, github, monkeypatch
):
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    FakeModel(_insights_answer).install(monkeypatch)
    job = await _repo_job(db_session, test_user, output_format="pptx", slide_count=8)

    await _run_report(job, test_user)

    row = await _reload(db_session, RepoReportJob, job.id)
    assert row.status == "completed", row.error
    assert minio.content_types[row.file_path] == PPTX_MIME
    blob = minio.objects[row.file_path]
    text = _pptx_text(blob)
    for expected in (
        "Widget Report",
        "Architecture Overview",
        "Recent Commits",
        "abcdef1: Fix stride prefetch distance",
        "ada: 42 contributions",
        "#7: ISB never engages",
        "C: 90.0%",
    ):
        assert expected in text, expected
    pictures = [
        shape
        for slide in _pptx(blob).slides
        for shape in slide.shapes
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE
    ]
    assert len(pictures) == 1  # the architecture diagram


async def test_repo_presentation_diagram_stays_on_the_configured_renderer(
    db_session, test_user, task_sessions, redis_fake, minio, github, monkeypatch
):
    import inspect

    from app.services import mermaid_renderer

    monkeypatch.setattr(settings, "KROKI_URL", "http://kroki-mermaid:8000")
    monkeypatch.setattr(settings, "KROKI_USE_FALLBACK", False)
    # A fresh singleton, and the render recorded at the renderer's own seam:
    # which base URL it chose is the claim, and it does not depend on how an
    # HTTP client was patched or which earlier test built the singleton.
    monkeypatch.setattr(mermaid_renderer, "_renderer", None)
    real = mermaid_renderer.MermaidRenderer._render_via_kroki
    chosen = []

    async def render_via_kroki(*args, **kwargs):
        bound = inspect.signature(real).bind(*args, **kwargs)
        chosen.append(bound.arguments["base_url"])
        from PIL import Image

        png = io.BytesIO()
        Image.new("RGB", (1, 1)).save(png, format="PNG")
        return png.getvalue()

    monkeypatch.setattr(
        mermaid_renderer.MermaidRenderer, "_render_via_kroki", render_via_kroki
    )
    FakeModel(_insights_answer).install(monkeypatch)
    job = await _repo_job(db_session, test_user, output_format="pptx")

    await _run_report(job, test_user)

    row = await _reload(db_session, RepoReportJob, job.id)
    assert row.status == "completed", row.error
    assert chosen == ["http://kroki-mermaid:8000"]


async def test_repo_report_docx_applies_the_requested_theme(
    db_session, test_user, task_sessions, redis_fake, minio, github, monkeypatch
):
    import docx
    from docx.shared import RGBColor

    from app.schemas.repo_report import ThemeColors, ThemeConfig

    FakeModel(_insights_answer).install(monkeypatch)
    theme = ThemeConfig(colors=ThemeColors(text_color="#aa0000")).model_dump()
    job = await _repo_job(db_session, test_user, custom_theme=theme)

    await _run_report(job, test_user)

    row = await _reload(db_session, RepoReportJob, job.id)
    assert row.status == "completed", row.error
    document = docx.Document(io.BytesIO(minio.objects[row.file_path]))
    assert document.styles["Normal"].font.color.rgb == RGBColor(0xAA, 0x00, 0x00)


@pytest.mark.parametrize(
    "overrides,cause",
    [
        (
            {"adhoc_url": "https://github.com/acme/missing"},
            "Unexpected error: Failed to initialize GitHub connector for "
            "acme/missing. Token provided: False.",
        ),
        ({"adhoc_url": None}, "No source_id or adhoc_url provided"),
    ],
    ids=["repository-not-found", "no-repository"],
)
async def test_repo_report_that_cannot_analyse_fails_naming_the_cause(
    db_session,
    test_user,
    task_sessions,
    redis_fake,
    minio,
    github,
    monkeypatch,
    overrides,
    cause,
):
    model = FakeModel(_insights_answer).install(monkeypatch)
    job = await _repo_job(db_session, test_user, **overrides)

    await _run_report(job, test_user)

    row = await _reload(db_session, RepoReportJob, job.id)
    assert row.status == "failed"
    assert row.error.startswith(cause)
    assert row.completed_at is not None
    assert row.file_path is None and minio.objects == {}
    assert model.calls == []
    last = redis_fake.on(f"repo_report:{job.id}:progress")[-1]
    assert last["status"] == "failed"
    assert last["error"] in row.error


async def test_repo_report_whose_upload_fails_is_failed_with_the_cause(
    db_session, test_user, task_sessions, redis_fake, github, monkeypatch
):
    FakeMinio(fail_uploads=ConnectionError("minio:9000 refused")).install(monkeypatch)
    FakeModel(_insights_answer).install(monkeypatch)
    job = await _repo_job(db_session, test_user)

    await _run_report(job, test_user)

    row = await _reload(db_session, RepoReportJob, job.id)
    assert row.status == "failed"
    assert row.error == "Unexpected error: minio:9000 refused"
    assert row.file_path is None and row.file_size is None


async def test_missing_repo_report_job_is_a_no_op(
    db_session, test_user, task_sessions, redis_fake, minio, github
):
    await repo_report_tasks._generate_repo_report_async(str(uuid4()), str(test_user.id))
    assert redis_fake.messages == [] and github.requests == []


async def test_cancelled_repo_report_job_is_left_alone(
    db_session, test_user, task_sessions, redis_fake, minio, github
):
    job = await _repo_job(db_session, test_user, status="cancelled")

    await _run_report(job, test_user)

    row = await _reload(db_session, RepoReportJob, job.id)
    assert row.status == "cancelled" and row.started_at is None
    assert redis_fake.messages == [] and github.requests == []


def test_generate_repo_report_task_run_completes_a_job(
    file_db, redis_fake, minio, github, monkeypatch
):
    FakeModel(_insights_answer).install(monkeypatch)
    user = _seed_user(file_db)
    job = file_db(lambda db: _repo_job(db, user))

    repo_report_tasks.generate_repo_report_task.run(str(job.id), str(user.id))

    row = _get(file_db, RepoReportJob, job.id)
    assert row.status == "completed", row.error
    assert row.file_size == len(minio.objects[row.file_path])


def test_generate_repo_report_task_run_on_a_missing_job(
    file_db, redis_fake, minio, github
):
    repo_report_tasks.generate_repo_report_task.run(str(uuid4()), str(uuid4()))
    assert redis_fake.messages == [] and github.requests == []
