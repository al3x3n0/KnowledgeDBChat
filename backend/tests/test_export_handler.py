"""`export_document`, called through its real handler.

The earlier version of this file copied the handler's markdown-to-slides loop
into the test and asserted on the copy, and "validated" parameters against
literals it wrote itself. These tests call the handler registered in the
provider. DOCX and PDF are built for real and opened again; `pptx` is stubbed
in conftest, so for PPTX the builder is replaced and the outline handed to it
is what gets checked. The TeX engine is the other replaced edge.
"""

import io
import re
from types import SimpleNamespace
from uuid import uuid4

import pytest
from sqlalchemy import select

from app.models.document import Document
from app.services import pptx_builder
from app.services.agent_tool_dispatch import (
    AgentToolExecutionContext,
    build_autonomous_document_authoring_provider,
)
from app.services.docx_builder import DOCXBuilder, markdown_to_content_items
from app.services.latex_compiler_service import LatexCompileResult, LatexCompilerService
from app.services.pdf_builder import PDFBuilder
from app.services.storage_service import storage_service

pytestmark = pytest.mark.unit

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
PPTX_MIME = "application/vnd.openxmlformats-officedocument.presentationml.presentation"

SAMPLE_MARKDOWN = """\
# Introduction

This is the introduction paragraph.

## Methods

- Method A
- Method B
- Method C

## Results

The results are shown below.

```python
print("hello")
```

## Conclusion

In conclusion, everything works.
"""


def _workspace(markdown=SAMPLE_MARKDOWN, title="Cache Study"):
    return {"plan": {"title": title, "sections": []}, "assembled_markdown": markdown}


@pytest.fixture
def job():
    return SimpleNamespace(id=uuid4(), user_id=uuid4(), name="author", config={})


@pytest.fixture
def export(job):
    """Call export_document the way the dispatcher does."""
    provider = build_autonomous_document_authoring_provider(SimpleNamespace())

    async def _export(params, state, db=None):
        return await provider._handlers["export_document"](
            params,
            AgentToolExecutionContext(
                mode="autonomous",
                db=db,
                service=None,
                user_id=str(job.user_id),
                job=job,
                state=state,
            ),
        )

    return _export


@pytest.fixture
def built(monkeypatch):
    """Capture the bytes the real DOCX and PDF builders return.

    The handler does not hand the file back, so the only way to look at what
    was built is to watch the real builder produce it.
    """
    captured = {}

    def _watch(cls, key):
        real = cls.build

        def build(self, *args, **kwargs):
            captured[key] = real(self, *args, **kwargs)
            return captured[key]

        monkeypatch.setattr(cls, "build", build)

    _watch(DOCXBuilder, "docx")
    _watch(PDFBuilder, "pdf")
    return captured


@pytest.fixture
def pptx(monkeypatch):
    """Replace the PPTX builder (python-pptx is stubbed) and keep its outline."""
    seen = SimpleNamespace(outlines=[], payload=b"PK\x03\x04fake-pptx")

    class FakePPTXBuilder:
        def build(self, outline=None, **kwargs):
            seen.outlines.append(outline)
            return seen.payload

    monkeypatch.setattr(pptx_builder, "PPTXBuilder", FakePPTXBuilder)
    return seen


@pytest.fixture
def uploads(monkeypatch):
    """Object storage, recording anything stored."""
    stored = []

    async def initialize():
        return None

    async def upload_to_path(object_path, content, content_type=None):
        stored.append(
            {"path": object_path, "content": content, "content_type": content_type}
        )
        return object_path

    async def get_presigned_download_url(object_path, expiry=None):
        return f"https://storage.test/{object_path}?signed=1"

    monkeypatch.setattr(storage_service, "initialize", initialize)
    monkeypatch.setattr(storage_service, "upload_to_path", upload_to_path)
    monkeypatch.setattr(
        storage_service, "get_presigned_download_url", get_presigned_download_url
    )
    return stored


def _docx_paragraphs(data):
    from docx import Document as open_docx

    return list(open_docx(io.BytesIO(data)).paragraphs)


# ---------------------------------------------------------------------------
# The markdown parser the DOCX and PDF paths share (real code, kept)
# ---------------------------------------------------------------------------


class TestMarkdownToContentItems:
    """Verify markdown_to_content_items produces expected structure."""

    def test_parses_headings(self):
        items = markdown_to_content_items("# Title\n\n## Subtitle\n\nText here.")
        headings = [i for i in items if i["type"] == "heading"]
        assert len(headings) >= 2
        assert headings[0]["level"] == 1
        assert headings[0]["text"] == "Title"
        assert headings[1]["level"] == 2
        assert headings[1]["text"] == "Subtitle"

    def test_parses_bullet_lists(self):
        md = "- Item 1\n- Item 2\n- Item 3"
        items = markdown_to_content_items(md)
        bullets = [i for i in items if i["type"] == "bullet_list"]
        assert len(bullets) == 1
        assert bullets[0]["items"] == ["Item 1", "Item 2", "Item 3"]

    def test_parses_code_blocks(self):
        md = "```python\nprint('hi')\n```"
        items = markdown_to_content_items(md)
        code = [i for i in items if i["type"] == "code_block"]
        assert len(code) == 1
        assert "print" in code[0]["code"]
        assert code[0]["language"] == "python"

    def test_empty_markdown(self):
        items = markdown_to_content_items("")
        assert items == []

    def test_full_document(self):
        items = markdown_to_content_items(SAMPLE_MARKDOWN)
        types = [i["type"] for i in items]
        assert "heading" in types
        assert "bullet_list" in types
        assert "code_block" in types


# ---------------------------------------------------------------------------
# Refusals
# ---------------------------------------------------------------------------


class TestExportRefusals:
    @pytest.mark.parametrize(
        "state",
        [
            {},
            {"document_workspace": {}},
            {"document_workspace": {"plan": {"title": "T"}}},
            {"document_workspace": {"plan": {"title": "T"}, "assembled_markdown": ""}},
        ],
    )
    async def test_nothing_assembled_means_nothing_to_export(
        self, export, built, state
    ):
        result = await export({"format": "docx"}, state)

        assert result == {
            "error": "No assembled document. Use assemble_document first."
        }
        assert built == {}

    @pytest.mark.parametrize("fmt", ["html", "", "odt", None])
    async def test_a_format_the_tool_does_not_offer_is_refused(
        self, export, built, fmt
    ):
        state = {"document_workspace": _workspace()}
        params = {} if fmt is None else {"format": fmt}

        result = await export(params, state)

        assert result["error"].startswith("Unsupported format")
        assert "docx, pdf, pptx, or latex" in result["error"]
        assert built == {}
        assert "export_artifacts" not in state["document_workspace"]

    async def test_the_refusal_covers_exactly_the_formats_the_spec_advertises(
        self, export
    ):
        from app.services.agent_tools import get_tool_by_name

        spec = get_tool_by_name("export_document")["parameters"]
        advertised = spec["properties"]["format"]["enum"]
        assert spec["required"] == ["format"]

        result = await export({"format": "html"}, {"document_workspace": _workspace()})

        named = set(re.findall(r"docx|pdf|pptx|latex", result["error"]))
        assert named == set(advertised)

    async def test_a_document_over_the_size_cap_is_refused_before_building(
        self, export, pptx
    ):
        state = {"document_workspace": _workspace("x" * 500_001)}

        result = await export({"format": "pptx"}, state)

        assert "too large" in result["error"]
        assert "500001" in result["error"]
        assert pptx.outlines == []
        assert "export_artifacts" not in state["document_workspace"]

    async def test_a_document_exactly_at_the_cap_is_exported(self, export, pptx):
        state = {"document_workspace": _workspace("x" * 500_000)}

        result = await export({"format": "pptx"}, state)

        assert result.get("success") is True, result
        assert len(pptx.outlines) == 1

    async def test_a_builder_that_fails_reports_it_and_records_no_artifact(
        self, export, monkeypatch
    ):
        def explode(self, *args, **kwargs):
            raise RuntimeError("template missing")

        monkeypatch.setattr(DOCXBuilder, "build", explode)
        state = {"document_workspace": _workspace()}

        result = await export({"format": "docx"}, state)

        assert "success" not in result
        assert "template missing" in result["error"]
        assert not state["document_workspace"].get("export_artifacts")


# ---------------------------------------------------------------------------
# DOCX — built for real and opened again
# ---------------------------------------------------------------------------


class TestExportDocx:
    async def test_the_export_is_a_real_word_document(self, export, built):
        state = {"document_workspace": _workspace()}

        result = await export({"format": "docx"}, state)

        assert result.get("success") is True, result
        assert built["docx"].startswith(b"PK")
        assert _docx_paragraphs(built["docx"])
        assert result["data"]["type"] == "exported_document"
        assert result["data"]["format"] == "docx"
        assert result["data"]["title"] == "Cache Study"
        assert result["data"]["mime_type"] == DOCX_MIME
        assert result["data"]["size_bytes"] == len(built["docx"])

    async def test_headings_from_the_markdown_are_headings_in_the_document(
        self, export, built
    ):
        await export({"format": "docx"}, {"document_workspace": _workspace()})

        headings = {
            p.text: p.style.name
            for p in _docx_paragraphs(built["docx"])
            if p.style.name.startswith("Heading")
        }
        for expected in ("Introduction", "Methods", "Results", "Conclusion"):
            assert expected in headings, headings
        assert headings["Introduction"] == "Heading 1"
        assert headings["Methods"] == "Heading 2"

    async def test_bullets_paragraphs_and_code_all_reach_the_document(
        self, export, built
    ):
        await export({"format": "docx"}, {"document_workspace": _workspace()})

        paragraphs = _docx_paragraphs(built["docx"])
        text = "\n".join(p.text for p in paragraphs)
        bullets = [p.text for p in paragraphs if "Bullet" in p.style.name]
        assert bullets == ["Method A", "Method B", "Method C"]
        assert "This is the introduction paragraph." in text
        assert "The results are shown below." in text
        assert 'print("hello")' in text
        assert "In conclusion, everything works." in text
        assert "```" not in text, "the code fence was written out literally"
        assert "## " not in text, "heading markup was written out literally"

    async def test_the_plan_title_is_the_document_title(self, export, built):
        await export(
            {"format": "docx"},
            {"document_workspace": _workspace(title="L2 Prefetcher Report")},
        )

        from docx import Document as open_docx

        document = open_docx(io.BytesIO(built["docx"]))
        all_text = "\n".join(p.text for p in document.paragraphs)
        assert (
            document.core_properties.title == "L2 Prefetcher Report"
            or "L2 Prefetcher Report" in all_text
        )

    async def test_format_is_matched_whatever_its_case(self, export, built):
        result = await export(
            {"format": "  DOCX "}, {"document_workspace": _workspace()}
        )

        assert result.get("success") is True, result
        assert result["data"]["format"] == "docx"
        assert "docx" in built

    async def test_each_export_is_recorded_on_the_workspace(self, export, built):
        state = {"document_workspace": _workspace()}

        first = await export({"format": "docx"}, state)
        second = await export({"format": "pdf"}, state)

        recorded = state["document_workspace"]["export_artifacts"]
        assert [a["format"] for a in recorded] == ["docx", "pdf"]
        assert recorded[0] == first["data"]
        assert recorded[1] == second["data"]


# ---------------------------------------------------------------------------
# PDF — built for real and read back
# ---------------------------------------------------------------------------


class TestExportPdf:
    async def test_the_export_is_a_real_pdf(self, export, built):
        state = {"document_workspace": _workspace()}

        result = await export({"format": "pdf"}, state)

        assert result.get("success") is True, result
        assert built["pdf"].startswith(b"%PDF")
        assert b"%%EOF" in built["pdf"][-1024:]
        assert result["data"]["format"] == "pdf"
        assert result["data"]["mime_type"] == "application/pdf"
        assert result["data"]["size_bytes"] == len(built["pdf"])

    async def test_the_markdown_content_is_in_the_pdf(self, export, built):
        pypdf = pytest.importorskip("pypdf")

        await export({"format": "pdf"}, {"document_workspace": _workspace()})

        reader = pypdf.PdfReader(io.BytesIO(built["pdf"]))
        text = "\n".join(page.extract_text() for page in reader.pages)
        for expected in (
            "Introduction",
            "Methods",
            "Method A",
            "Method C",
            "The results are shown below.",
            "In conclusion, everything works.",
        ):
            assert expected in text, text
        assert "## " not in text, "heading markup was written out literally"


# ---------------------------------------------------------------------------
# PPTX — python-pptx is stubbed, so the outline is what is checked
# ---------------------------------------------------------------------------


class TestExportPptx:
    async def test_the_builder_is_handed_one_slide_per_section(self, export, pptx):
        state = {"document_workspace": _workspace()}

        result = await export({"format": "pptx"}, state)

        assert result.get("success") is True, result
        assert len(pptx.outlines) == 1
        outline = pptx.outlines[0]
        assert outline.title == "Cache Study"
        assert [s.title for s in outline.slides] == [
            "Introduction",
            "Methods",
            "Results",
            "Conclusion",
        ]
        assert [s.slide_number for s in outline.slides] == [1, 2, 3, 4]
        assert result["data"]["format"] == "pptx"
        assert result["data"]["mime_type"] == PPTX_MIME
        assert result["data"]["size_bytes"] == len(pptx.payload)

    async def test_bullets_and_prose_become_slide_content(self, export, pptx):
        await export({"format": "pptx"}, {"document_workspace": _workspace()})

        slides = {s.title: s for s in pptx.outlines[0].slides}
        assert slides["Methods"].content == ["Method A", "Method B", "Method C"]
        assert slides["Introduction"].content == ["This is the introduction paragraph."]
        assert "In conclusion, everything works." in slides["Conclusion"].content

    @pytest.mark.parametrize("marker", ["-", "*", "+"])
    async def test_every_markdown_bullet_marker_is_stripped(self, export, pptx, marker):
        markdown = f"## Options\n\n{marker} first\n{marker} second\n"

        await export({"format": "pptx"}, {"document_workspace": _workspace(markdown)})

        assert pptx.outlines[0].slides[0].content == ["first", "second"]

    async def test_no_slide_carries_more_than_ten_bullets(self, export, pptx):
        markdown = "## Big Section\n\n" + "\n".join(f"- Item {i}" for i in range(20))

        await export({"format": "pptx"}, {"document_workspace": _workspace(markdown)})

        assert all(len(s.content) <= 10 for s in pptx.outlines[0].slides)

    @pytest.mark.xfail(
        strict=True,
        reason=(
            "The PPTX path keeps `bullets[:10]` per section and discards the "
            "rest: a section with 20 bullets exports 10 of them, with no "
            "continuation slide and nothing in the result saying content was "
            "dropped."
        ),
    )
    async def test_a_long_section_loses_no_content(self, export, pptx):
        markdown = "## Big Section\n\n" + "\n".join(f"- Item {i}" for i in range(20))

        await export({"format": "pptx"}, {"document_workspace": _workspace(markdown)})

        exported = [line for s in pptx.outlines[0].slides for line in s.content]
        assert exported == [f"Item {i}" for i in range(20)]

    @pytest.mark.xfail(
        strict=True,
        reason=(
            "The PPTX path treats every non-heading line as a bullet, so the "
            "fence lines of a code block ('```python', '```') are exported "
            "as slide bullets. The DOCX/PDF paths parse the fence."
        ),
    )
    async def test_code_fence_markers_are_not_slide_bullets(self, export, pptx):
        await export({"format": "pptx"}, {"document_workspace": _workspace()})

        results = [s for s in pptx.outlines[0].slides if s.title == "Results"][0]
        assert not any(line.startswith("```") for line in results.content)
        assert 'print("hello")' in results.content

    async def test_markdown_with_no_content_still_yields_a_title_slide(
        self, export, pptx
    ):
        state = {"document_workspace": _workspace("  \n\n  ", title="Fallback")}

        result = await export({"format": "pptx"}, state)

        assert result.get("success") is True, result
        slides = pptx.outlines[0].slides
        assert len(slides) == 1
        assert slides[0].slide_type == "title"
        assert slides[0].title == "Fallback"

    async def test_a_document_without_h2_headings_is_one_slide(self, export, pptx):
        markdown = "# Only H1\n\nSome content without H2 headings."

        await export({"format": "pptx"}, {"document_workspace": _workspace(markdown)})

        slides = pptx.outlines[0].slides
        assert [s.title for s in slides] == ["Only H1"]
        assert slides[0].content == ["Some content without H2 headings."]


# ---------------------------------------------------------------------------
# LaTeX — the TeX engine is the replaced edge
# ---------------------------------------------------------------------------


class TestExportLatex:
    @pytest.mark.xfail(
        strict=True,
        reason=(
            "export_document calls LatexCompilerService.compile_to_pdf(...) "
            "on the class, but it is an instance method: every latex export "
            "fails with \"missing 1 required positional argument: 'self'\" "
            "and is reported as 'Export failed'. The advertised latex format "
            "has never worked."
        ),
    )
    async def test_latex_export_compiles_and_returns_the_pdf(self, export, monkeypatch):
        compiled = []

        def compile_to_pdf(self, *, tex_source, timeout_seconds, max_source_chars, **_):
            compiled.append(tex_source)
            return LatexCompileResult(
                success=True,
                engine="tectonic",
                pdf_bytes=b"%PDF-1.5 compiled",
                log="",
                violations=[],
            )

        monkeypatch.setattr(LatexCompilerService, "compile_to_pdf", compile_to_pdf)
        state = {"document_workspace": _workspace()}

        result = await export({"format": "latex"}, state)

        assert result.get("success") is True, result
        assert len(compiled) == 1
        assert result["data"]["format"] == "latex"
        assert result["data"]["mime_type"] == "application/pdf"
        assert result["data"]["size_bytes"] == len(b"%PDF-1.5 compiled")

    async def test_a_compile_failure_is_reported_with_the_log(
        self, export, monkeypatch
    ):
        def compile_to_pdf(*args, **kwargs):
            return LatexCompileResult(
                success=False,
                engine=None,
                pdf_bytes=None,
                log="! Undefined control sequence.",
                violations=[],
            )

        monkeypatch.setattr(LatexCompilerService, "compile_to_pdf", compile_to_pdf)
        state = {"document_workspace": _workspace()}

        result = await export({"format": "latex"}, state)

        assert "success" not in result
        assert "LaTeX compilation failed" in result["error"]
        assert "Undefined control sequence" in result["error"]
        assert not state["document_workspace"].get("export_artifacts")


# ---------------------------------------------------------------------------
# What the export leaves behind
# ---------------------------------------------------------------------------


class TestExportIsDelivered:
    @pytest.mark.xfail(
        strict=True,
        reason=(
            "export_document builds the file and discards it: the bytes are "
            "never uploaded to storage, never written anywhere and not "
            "returned. The result carries only a size and a MIME type, so "
            "there is nothing a person or a later tool can download."
        ),
    )
    async def test_the_exported_file_can_be_retrieved(self, export, built, uploads):
        state = {"document_workspace": _workspace()}

        result = await export({"format": "docx"}, state)

        assert result.get("success") is True, result
        stored = [u for u in uploads if u["content"] == built["docx"]]
        assert stored, "the built DOCX was never stored"
        assert stored[0]["content_type"] == DOCX_MIME
        assert result["data"].get("url") or result["data"].get("object_path")

    async def test_nothing_is_saved_to_the_kb_unless_asked(
        self, export, built, db_session
    ):
        state = {"document_workspace": _workspace()}

        result = await export({"format": "docx"}, state, db=db_session)

        assert result.get("success") is True, result
        assert "document_id" not in result["data"]
        assert (await db_session.execute(select(Document))).scalars().all() == []

    @pytest.mark.xfail(
        strict=True,
        reason=(
            "persist_to_kb builds Document(...) without source_id and "
            "source_identifier, both NOT NULL, so the flush fails. The "
            "failure is caught and logged, and the tool still answers "
            "success -- with no document_id and nothing in the knowledge "
            "base."
        ),
    )
    async def test_persist_to_kb_stores_a_document(self, export, built, db_session):
        state = {"document_workspace": _workspace()}

        try:
            result = await export(
                {"format": "docx", "persist_to_kb": True}, state, db=db_session
            )

            assert result.get("success") is True, result
            document_id = result["data"].get("document_id")
            assert document_id, "export succeeded but nothing was saved to the KB"
            stored = (await db_session.execute(select(Document))).scalars().all()
            assert [str(d.id) for d in stored] == [document_id]
            assert stored[0].title.startswith("Cache Study")
            assert "Method A" in stored[0].content
            assert stored[0].extra_metadata["format"] == "docx"
        finally:
            await db_session.rollback()

    @pytest.mark.xfail(
        strict=True,
        reason=(
            "The export_document spec declares `latex_project_id` ('Existing "
            "LaTeX project to export into') and the handler never reads it: "
            "the name does not appear in the document-authoring provider."
        ),
    )
    def test_latex_project_id_is_read_by_the_handler(self):
        import inspect

        from app.services import agent_tool_dispatch
        from app.services.agent_tools import get_tool_by_name

        source = inspect.getsource(
            agent_tool_dispatch.build_autonomous_document_authoring_provider
        )
        declared = get_tool_by_name("export_document")["parameters"]["properties"]
        assert "latex_project_id" in declared
        for param in declared:
            assert f'"{param}"' in source, f"export_document never reads {param}"


# ---------------------------------------------------------------------------
# Builder signatures the handler depends on (real code, kept)
# ---------------------------------------------------------------------------


class TestBuilderSignatures:
    """The handler calls `build(title=..., content_items=...)` / `build(outline=...)`."""

    @pytest.mark.parametrize("builder_cls", [DOCXBuilder, PDFBuilder])
    def test_document_builders_take_title_and_content_items(self, builder_cls):
        import inspect

        params = inspect.signature(builder_cls.build).parameters
        assert "title" in params
        assert "content_items" in params
        assert not hasattr(builder_cls, "build_from_markdown")

    def test_pptx_builder_takes_an_outline(self):
        import inspect

        params = inspect.signature(pptx_builder.PPTXBuilder.build).parameters
        assert "outline" in params
        assert not hasattr(pptx_builder.PPTXBuilder, "build_from_markdown")
